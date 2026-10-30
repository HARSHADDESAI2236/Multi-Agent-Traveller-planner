"""Small, testable helpers for TravelMate's additive travel-planning features."""
from __future__ import annotations

import hashlib
import json
from dataclasses import asdict, is_dataclass
from datetime import date
from typing import Any, Iterable
from urllib.parse import quote_plus


def google_maps_search_url(hotel_name: str, destination: str = "") -> str:
    """Return a Google Maps search URL, never a fake hotel detail URL."""
    name = (hotel_name or "").strip()
    place = (destination or "").strip()
    if not name:
        raise ValueError("Hotel name is required")
    query = " ".join(part for part in (name, place) if part)
    return "https://www.google.com/maps/search/?api=1&query=" + quote_plus(query)


def google_hotel_photos_url(hotel_name: str, destination: str = "") -> str:
    """Search Google for photos/reviews; provider results remain authoritative."""
    name = (hotel_name or "").strip()
    if not name:
        raise ValueError("Hotel name is required")
    query = " ".join(part for part in (name, destination.strip(), "photos reviews") if part)
    return "https://www.google.com/search?q=" + quote_plus(query)


def google_location_search_url(query: str) -> str:
    """Build a safe Google Maps search URL for any user-entered location."""
    clean_query = (query or "").strip()
    if not clean_query:
        raise ValueError("Enter a destination to search")
    return "https://www.google.com/maps/search/?api=1&query=" + quote_plus(clean_query)


def search_destination_catalog(categories: dict, query: str, limit: int = 6) -> list[dict]:
    """Search a nested destination catalogue safely, case-insensitively, with a bounded result count."""
    needle = (query or "").strip().casefold()
    if not needle:
        return []
    safe_limit = max(1, min(int(limit), 10))
    matches = []
    for category in categories.values():
        for place in category.get("places", []):
            searchable = " ".join([
                str(place.get("name", "")), str(place.get("description", "")),
                *(str(item) for item in place.get("highlights", [])),
            ]).casefold()
            if needle in searchable:
                matches.append(place)
                if len(matches) >= safe_limit:
                    return matches
    return matches


def trip_fingerprint(origin: str, destination: str, start_date: Any, end_date: Any,
                     budget: float, travelers: int, preferences: Iterable[str],
                     pace: str = "Balanced", accommodation: str = "Any") -> str:
    """Stable idempotency key to avoid re-running agents on rapid repeat submits."""
    payload = {
        "origin": origin.strip().casefold(),
        "destination": destination.strip().casefold(),
        "start_date": str(start_date),
        "end_date": str(end_date),
        "budget": round(float(budget), 2),
        "travelers": int(travelers),
        "preferences": sorted(str(p).strip().casefold() for p in preferences),
        "pace": pace,
        "accommodation": accommodation,
    }
    return hashlib.sha256(json.dumps(payload, sort_keys=True).encode("utf-8")).hexdigest()


def validate_trip(origin: str, destination: str, start_date: date, end_date: date,
                  budget: float, travelers: int) -> list[str]:
    errors = []
    if not (origin or "").strip():
        errors.append("Please enter your starting location.")
    if not (destination or "").strip():
        errors.append("Please enter your destination.")
    if start_date >= end_date:
        errors.append("Your return date must be after your departure date.")
    if budget <= 0:
        errors.append("Budget must be greater than zero.")
    if travelers < 1 or travelers > 50:
        errors.append("Travelers must be between 1 and 50.")
    return errors


def _plain(value: Any) -> Any:
    if is_dataclass(value):
        return asdict(value)
    if hasattr(value, "__dict__"):
        return {k: _plain(v) for k, v in vars(value).items() if not k.startswith("_")}
    if isinstance(value, (list, tuple)):
        return [_plain(v) for v in value]
    if isinstance(value, dict):
        return {str(k): _plain(v) for k, v in value.items()}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def trip_to_dict(plan: Any) -> dict[str, Any]:
    """Convert the project's dataclass plan into export-friendly JSON data."""
    return _plain(plan)


def trip_to_text(plan: Any) -> str:
    """Create a readable, dependency-free itinerary summary."""
    request = getattr(plan, "request", None)
    lines = ["TRAVELMATE — TRIP PLAN", "=" * 32]
    if request:
        lines += [f"Route: {request.origin} → {request.destination}",
                  f"Dates: {request.start_date} to {request.end_date}",
                  f"Travelers: {request.travelers}", f"Budget: ₹{request.budget:,.0f}", ""]
    lines.append("FLIGHT OPTIONS")
    for f in getattr(plan, "flights", []):
        lines.append(f"- {f.airline}: {f.departure} → {f.arrival}; {f.duration_hours}h; ₹{f.price:,.0f}")
    lines += ["", "HOTEL RECOMMENDATIONS"]
    for h in getattr(plan, "hotels", []):
        lines.append(f"- {h.name}, {h.neighborhood}; rating {h.rating}; ₹{h.price_per_night:,.0f}/night")
    lines += ["", "WEATHER"]
    for w in getattr(plan, "weather", []):
        lines.append(f"- {w.date}: {w.summary}, {w.high_c}° / {w.low_c}°C")
    lines += ["", "DAY-BY-DAY ITINERARY"]
    for day in getattr(plan, "itinerary", []):
        lines.append(f"Day {day.day_number} — {day.date}")
        lines.extend(f"  • {activity}" for activity in day.activities)
    lines += ["", "Prices and availability are estimates and must be verified with providers."]
    return "\n".join(lines)


def budget_breakdown(plan: Any, days: int | None = None) -> dict[str, float]:
    """Transparent indicative budget split based on current agent estimates."""
    request = getattr(plan, "request", None)
    budget = max(0.0, float(getattr(request, "budget", 0) or 0))
    travelers = max(1, int(getattr(request, "travelers", 1) or 1))
    if days is None and request:
        try:
            days = max(1, (date.fromisoformat(request.end_date) - date.fromisoformat(request.start_date)).days)
        except (ValueError, TypeError):
            days = 1
    days = max(1, int(days or 1))
    flight = min((float(f.price) for f in getattr(plan, "flights", [])), default=0.0) * travelers
    hotel = min((float(h.price_per_night) for h in getattr(plan, "hotels", [])), default=0.0) * days
    remaining = max(0.0, budget - flight - hotel)
    food = remaining * 0.45
    local_transport = remaining * 0.25
    activities = remaining * 0.20
    contingency = remaining * 0.10
    return {"Flights (estimate)": round(flight, 2), "Accommodation (estimate)": round(hotel, 2),
            "Food (suggested)": round(food, 2), "Local transport (suggested)": round(local_transport, 2),
            "Activities (suggested)": round(activities, 2), "Contingency (suggested)": round(contingency, 2),
            "Total budget": round(budget, 2), "Estimated / suggested total": round(flight + hotel + remaining, 2)}


def trip_to_pdf(plan: Any) -> bytes:
    """Render a PDF itinerary using ReportLab's flowable layout."""
    from io import BytesIO
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=18 * mm,
                                 leftMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    styles = getSampleStyleSheet()
    headings = {"FLIGHT OPTIONS", "HOTEL RECOMMENDATIONS", "WEATHER", "DAY-BY-DAY ITINERARY"}
    story = []
    for line in trip_to_text(plan).splitlines():
        if line.strip() == "=" * 32:
            continue
        safe_line = escape(line).replace("₹", "INR ") or "&nbsp;"
        style = styles["Title"] if line == "TRAVELMATE — TRIP PLAN" else styles["Heading2"] if line in headings else styles["BodyText"]
        story.append(Paragraph(safe_line, style))
        story.append(Spacer(1, 4))
    document.build(story)
    return buffer.getvalue()


def trip_to_docx(plan: Any) -> bytes:
    """Render a Word document itinerary."""
    from io import BytesIO
    from docx import Document

    document = Document()
    headings = {"FLIGHT OPTIONS", "HOTEL RECOMMENDATIONS", "WEATHER", "DAY-BY-DAY ITINERARY"}
    for line in trip_to_text(plan).splitlines():
        if line.strip() == "=" * 32:
            continue
        if line == "TRAVELMATE — TRIP PLAN":
            document.add_heading(line, level=0)
        elif line in headings:
            document.add_heading(line, level=1)
        elif line.startswith("- ") or line.startswith("  • "):
            document.add_paragraph(line.lstrip("- •"), style="List Bullet")
        elif line.strip():
            document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def get_or_create_cached_plan(cache: dict[str, Any], key: str, factory, max_entries: int = 3) -> tuple[Any, bool]:
    """Return (plan, reused). The caller owns/persists the returned cache dict."""
    if key in cache:
        return cache[key], True
    plan = factory()
    cache[key] = plan
    while len(cache) > max(1, int(max_entries)):
        cache.pop(next(iter(cache)))
    return plan, False
